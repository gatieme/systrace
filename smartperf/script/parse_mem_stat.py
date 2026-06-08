#!/usr/bin/env python3
"""Parse mem_stat_jq.txt (from dump_5s_hm_2.sh) and export memory stats to xlsx with line charts."""

from __future__ import annotations

import re
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

LOCAL_TZ = ZoneInfo("Asia/Shanghai")

TS_RE = re.compile(
    r"^([A-Z][a-z]{2} [A-Z][a-z]{2}\s+\d{1,2} \d{2}:\d{2}:\d{2} [A-Z]{3} \d{4})$"
)
MEMINFO_KB_RE = re.compile(r"^([A-Za-z0-9()]+):\s+(\d+)\s*kB\s*$")
MEMINFO_COUNT_RE = re.compile(r"^([A-Za-z0-9()]+):\s+(\d+)\s*$")

# Keys exported to the time-series sheet (from /proc/meminfo)
MEMINFO_KEYS = (
    "MemTotal",
    "MemFree",
    "MemAvailable",
    "Cached",
    "Active",
    "Inactive",
    "SwapTotal",
    "SwapFree",
    "ZramUsed",
    "AnonPages",
    "Committed_AS",
    "IonTotalUsed",
    "GpuTotalUsed",
    "DmaHeapTotalUsed",
    "RsvTotalUsed",
    "Slab",
    "Dirty",
)


def parse_timestamp(line: str) -> datetime | None:
    line = " ".join(line.strip().split())
    if not TS_RE.match(line):
        return None
    try:
        # %Z does not resolve "CST" on Windows; dump script uses CST.
        dt = datetime.strptime(line, "%a %b %d %H:%M:%S CST %Y")
    except ValueError:
        return None
    return dt.replace(tzinfo=LOCAL_TZ)


def parse_meminfo_block(lines: list[str]) -> dict[str, int]:
    values: dict[str, int] = {}
    in_meminfo = False
    for raw in lines:
        line = raw.rstrip("\n")
        if line == "cat /proc/meminfo":
            in_meminfo = True
            continue
        if in_meminfo and line.startswith("cat "):
            break
        if not in_meminfo:
            continue
        m = MEMINFO_KB_RE.match(line)
        if m:
            values[m.group(1)] = int(m.group(2))
            continue
        m = MEMINFO_COUNT_RE.match(line)
        if m:
            values[m.group(1)] = int(m.group(2))
    return values


def split_snapshots(text: str) -> list[tuple[datetime, list[str]]]:
    snapshots: list[tuple[datetime, list[str]]] = []
    current_ts: datetime | None = None
    current_lines: list[str] = []

    for line in text.splitlines():
        ts = parse_timestamp(line)
        if ts is not None:
            if current_ts is not None:
                snapshots.append((current_ts, current_lines))
            current_ts = ts
            current_lines = []
            continue
        if current_ts is not None:
            current_lines.append(line)

    if current_ts is not None:
        snapshots.append((current_ts, current_lines))
    return snapshots


def build_dataframe(snapshots: list[tuple[datetime, list[str]]]) -> pd.DataFrame:
    rows: list[dict] = []
    t0 = snapshots[0][0] if snapshots else None

    for ts, block_lines in snapshots:
        mem = parse_meminfo_block(block_lines)
        row: dict = {
            "timestamp": ts,
            "timestamp_str": ts.strftime("%Y-%m-%d %H:%M:%S"),
            "elapsed_sec": round((ts - t0).total_seconds(), 1) if t0 else 0.0,
        }
        for key in MEMINFO_KEYS:
            row[key] = mem.get(key)

        mem_total = mem.get("MemTotal")
        mem_avail = mem.get("MemAvailable")
        mem_free = mem.get("MemFree")
        if mem_total is not None and mem_avail is not None:
            row["MemUsed"] = mem_total - mem_avail
        if mem_total is not None and mem_free is not None:
            row["MemUsedSimple"] = mem_total - mem_free
        swap_total = mem.get("SwapTotal")
        swap_free = mem.get("SwapFree")
        if swap_total is not None and swap_free is not None:
            row["SwapUsed"] = swap_total - swap_free

        rows.append(row)

    df = pd.DataFrame(rows)
    if df.empty:
        return df

    for col in (*MEMINFO_KEYS, "MemUsed", "MemUsedSimple", "SwapUsed"):
        if col in df.columns:
            df[f"{col}_MB"] = (df[col] / 1024).round(2)

    if "MemTotal" in df.columns and "MemUsed" in df.columns:
        df["MemUsed_pct"] = (df["MemUsed"] / df["MemTotal"] * 100).round(2)

    return df


def export_xlsx(df: pd.DataFrame, output_path: Path, source: Path, interval_sec: int = 5) -> None:
    export_cols = [
        "timestamp_str",
        "elapsed_sec",
        "MemTotal_MB",
        "MemFree_MB",
        "MemAvailable_MB",
        "MemUsed_MB",
        "MemUsed_pct",
        "Cached_MB",
        "Active_MB",
        "Inactive_MB",
        "AnonPages_MB",
        "Committed_AS_MB",
        "GpuTotalUsed_MB",
        "IonTotalUsed_MB",
        "DmaHeapTotalUsed_MB",
        "SwapUsed_MB",
        "ZramUsed_MB",
        "Slab_MB",
        "Dirty_MB",
    ]
    export_cols = [c for c in export_cols if c in df.columns]
    ts_df = df[export_cols].copy()

    t_start = df["timestamp"].iloc[0]
    t_end = df["timestamp"].iloc[-1]
    duration_sec = (t_end - t_start).total_seconds() if len(df) > 1 else 0.0

    summary = pd.DataFrame(
        [
            {
                "source_file": str(source.name),
                "sample_count": len(df),
                "sample_interval_sec": interval_sec,
                "start_time": str(t_start),
                "end_time": str(t_end),
                "duration_sec": round(duration_sec, 1),
                "mem_used_mb_min": df["MemUsed_MB"].min() if "MemUsed_MB" in df else None,
                "mem_used_mb_max": df["MemUsed_MB"].max() if "MemUsed_MB" in df else None,
                "mem_used_mb_mean": round(df["MemUsed_MB"].mean(), 2) if "MemUsed_MB" in df else None,
                "mem_available_mb_min": df["MemAvailable_MB"].min()
                if "MemAvailable_MB" in df
                else None,
            }
        ]
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="summary", index=False)
        ts_df.to_excel(writer, sheet_name="mem_timeseries", index=False)

    _add_charts(output_path, ts_df)


LINE_COLORS = ("4472C4", "ED7D31", "70AD47", "FFC000")


def _style_series(ser: Series, color: str) -> None:
    ser.marker.symbol = "none"
    ser.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=color, w=25000)
    )


def _configure_axes(chart: ScatterChart, *, y_fmt: str, y_title: str, y_min: float, y_max: float, x_max: float) -> None:
    chart.varyColors = False
    chart.x_axis.title = "时间 (秒)"
    chart.y_axis.title = y_title
    chart.x_axis.tickLblPos = "nextTo"
    chart.y_axis.tickLblPos = "nextTo"
    chart.x_axis.majorTickMark = "out"
    chart.y_axis.majorTickMark = "out"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.numFmt = "0"
    chart.y_axis.numFmt = y_fmt
    chart.x_axis.scaling.min = 0
    chart.x_axis.scaling.max = x_max
    chart.y_axis.scaling.min = y_min
    chart.y_axis.scaling.max = y_max
    # 约每 30 秒一个横轴刻度，便于读出秒数
    chart.x_axis.majorUnit = 30 if x_max > 60 else 10


def _make_scatter_chart(
    ws,
    n_rows: int,
    elapsed_col: int,
    series_specs: list[tuple[str, int, str]],
    *,
    title: str,
    y_title: str,
    y_fmt: str,
    y_min: float,
    y_max: float,
    x_max: float,
) -> ScatterChart:
    chart = ScatterChart()
    chart.title = title
    chart.scatterStyle = "line"
    chart.height = 12
    chart.width = 24
    _configure_axes(chart, y_fmt=y_fmt, y_title=y_title, y_min=y_min, y_max=y_max, x_max=x_max)

    xvalues = Reference(ws, min_col=elapsed_col, min_row=2, max_row=n_rows)
    for idx, (series_title, value_col, color) in enumerate(series_specs):
        yvalues = Reference(ws, min_col=value_col, min_row=2, max_row=n_rows)
        ser = Series(yvalues, xvalues, title=series_title)
        _style_series(ser, color or LINE_COLORS[idx % len(LINE_COLORS)])
        chart.series.append(ser)

    return chart


def _add_charts(xlsx_path: Path, ts_df: pd.DataFrame) -> None:
    wb = load_workbook(xlsx_path)
    ws = wb["mem_timeseries"]
    n_rows = len(ts_df) + 1  # header + data

    if "elapsed_sec" not in ts_df.columns:
        wb.save(xlsx_path)
        return

    elapsed_col = list(ts_df.columns).index("elapsed_sec") + 1
    x_max = float(ts_df["elapsed_sec"].max())
    x_max = max(x_max, 5.0)

    def col_idx(name: str) -> int:
        return list(ts_df.columns).index(name) + 1

    chart_specs: list[tuple[str, str, str, list[str], str]] = [
        ("内存占用 (MemTotal - MemAvailable)", "内存占用 (MB)", "0", ["MemUsed_MB"], "MB"),
        ("可用/空闲内存", "内存 (MB)", "0", ["MemAvailable_MB", "MemFree_MB"], "MB"),
        ("GPU / ION 内存", "内存 (MB)", "0", ["GpuTotalUsed_MB", "IonTotalUsed_MB"], "MB"),
        ("缓存与匿名页", "内存 (MB)", "0", ["Cached_MB", "AnonPages_MB"], "MB"),
        ("内存占用率", "占用率 (%)", "0.0", ["MemUsed_pct"], "%"),
    ]

    chart_row = n_rows + 3
    for chart_title, y_title, y_fmt, value_cols, _ in chart_specs:
        cols_present = [c for c in value_cols if c in ts_df.columns]
        if not cols_present:
            continue

        y_min = float(ts_df[cols_present].min().min())
        y_max = float(ts_df[cols_present].max().max())
        if cols_present == ["MemUsed_pct"]:
            y_min = max(0.0, y_min - 1.0)
            y_max = min(100.0, y_max + 1.0)
        else:
            pad = max((y_max - y_min) * 0.05, 1.0)
            y_min = max(0.0, y_min - pad)
            y_max = y_max + pad

        series_specs = [
            (col.replace("_", " "), col_idx(col), LINE_COLORS[i % len(LINE_COLORS)])
            for i, col in enumerate(cols_present)
        ]
        chart = _make_scatter_chart(
            ws,
            n_rows,
            elapsed_col,
            series_specs,
            title=chart_title,
            y_title=y_title,
            y_fmt=y_fmt,
            y_min=y_min,
            y_max=y_max,
            x_max=x_max,
        )
        ws.add_chart(chart, f"A{chart_row}")
        chart_row += 20

    wb.save(xlsx_path)


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print("Usage: python parse_mem_stat.py <mem_stat_jq.txt> [output.xlsx]")
        return 1

    input_path = Path(argv[1]).resolve()
    if not input_path.exists():
        print(f"File not found: {input_path}")
        return 1

    output_path = (
        Path(argv[2]).resolve()
        if len(argv) > 2
        else input_path.with_suffix(".stats.xlsx")
    )

    print(f"Parsing {input_path} ...")
    text = input_path.read_text(encoding="utf-8", errors="replace")
    snapshots = split_snapshots(text)
    if not snapshots:
        print("No timestamp snapshots found.")
        return 1

    df = build_dataframe(snapshots)
    export_xlsx(df, output_path, input_path)

    print(f"Done. Output: {output_path}")
    print(f"  snapshots: {len(df)}")
    if "MemUsed_MB" in df.columns:
        print(
            f"  MemUsed_MB: min={df['MemUsed_MB'].min():.2f}, "
            f"max={df['MemUsed_MB'].max():.2f}, mean={df['MemUsed_MB'].mean():.2f}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
